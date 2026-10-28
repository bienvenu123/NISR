from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
CSV = ROOT / "outputs" / "eicv7_district_double_exclusion.csv"
OUTPUT = ROOT / "deliverables" / "rwanda_double_exclusion_policy_brief.docx"
NAVY = "17365D"
TEAL = "087E8B"
PALE = "EAF1F5"
GRAY = "59636E"


def set_cell_shading(cell, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def set_cell_border(cell, color="D9E0E5", size="5") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            margins.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def style_run(run, size=10, bold=False, color="263238", font="Arial"):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph(paragraph, before=0, after=4, line=1.04):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_para(doc, text, size=10, after=4, color="263238", bold_prefix=None):
    p = doc.add_paragraph()
    set_paragraph(p, after=after)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        style_run(r, size=size, bold=True, color=color)
        r = p.add_run(text[len(bold_prefix):])
        style_run(r, size=size, color=color)
    else:
        r = p.add_run(text)
        style_run(r, size=size, color=color)
    return p


def add_heading(doc, text, size=13):
    p = doc.add_paragraph()
    set_paragraph(p, before=5, after=4)
    r = p.add_run(text)
    style_run(r, size=size, bold=True, color=NAVY)
    return p


def format_table(table, header=True, widths=None, font_size=9.2):
    table.autofit = False
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            set_cell_margins(cell)
            if widths:
                cell.width = Inches(widths[col_index])
            if header and row_index == 0:
                set_cell_shading(cell, NAVY)
            elif row_index % 2 == 0:
                set_cell_shading(cell, PALE)
            for p in cell.paragraphs:
                set_paragraph(p, after=0, line=1.0)
                for run in p.runs:
                    style_run(run, size=font_size, bold=(header and row_index == 0),
                              color="FFFFFF" if header and row_index == 0 else "263238")


def build() -> None:
    df = pd.read_csv(CSV)
    poor_hh = df["estimated_poor_households"].sum()
    excluded_hh = df["estimated_double_excluded_poor_households"].sum()
    national_double = excluded_hh / poor_hh
    national_vup = (df["vup_coverage_among_poor"] * df["estimated_poor_households"]).sum() / poor_hh
    national_finance = (df["financial_access_among_poor"] * df["estimated_poor_households"]).sum() / poor_hh
    national_poverty = df["estimated_poor_population"].sum() / df["estimated_population"].sum()
    by_burden = df.sort_values("estimated_double_excluded_poor_households", ascending=False).head(5)

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.25)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string("263238")
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.04

    # Word's built-in Title style adds a blue bottom border in some themes.
    # Strip it so the brief follows the plain black heading treatment.
    title_style = doc.styles["Title"]
    title_ppr = title_style._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header.add_run("EVIDENCE BRIEF  |  RWANDA  |  EICV7 2023-24")
    style_run(r, size=8, bold=True, color="000000")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("District estimates are descriptive and should not be used to identify individual households")
    style_run(r, size=8, color=GRAY)

    p = doc.add_paragraph(style="Title")
    set_paragraph(p, after=2)
    r = p.add_run("Double exclusion in Rwanda")
    style_run(r, size=24, bold=True, color="000000")
    p = doc.add_paragraph()
    set_paragraph(p, after=8)
    r = p.add_run("Where poor households are missed by VUP and bank or mobile accounts")
    style_run(r, size=12, color="000000")

    add_para(doc,
        "About 17.5% of poor households are estimated to have neither an observed VUP participation flag nor a bank or mobile-money account. That is approximately 134,600 poor households nationally. The EICV7 data point to a combined access gap that can guide district-level investigation, while they do not show whether a household was eligible for VUP.",
        size=10.5, after=6)

    add_heading(doc, "National picture")
    metrics = [
        ("Population poverty rate", f"{national_poverty:.1%}", "Population weight; matches NISR's 27.4% national rate after rounding"),
        ("VUP coverage among poor households", f"{national_vup:.1%}", "VUP participation reported in EICV7 modules"),
        ("Bank or mobile account access among poor households", f"{national_finance:.1%}", "At least one household member reports an account"),
        ("Poor households with double exclusion", f"{national_double:.1%}", f"About {excluded_hh:,.0f} households, household weighted"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ["Measure", "Estimate", "Definition"]):
        cell.text = text
    for row in metrics:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = text
    format_table(table, widths=[2.2, 1.15, 3.4], font_size=8.7)

    add_heading(doc, "Districts with the largest estimated burden")
    add_para(doc, "The ranking uses the estimated number of poor households with double exclusion. Rates show the share of poor households affected.", size=9.3, after=4)
    table = doc.add_table(rows=1, cols=3)
    for cell, text in zip(table.rows[0].cells, ["District", "Double-exclusion rate", "Estimated households"]):
        cell.text = text
    for row in by_burden.itertuples(index=False):
        cells = table.add_row().cells
        cells[0].text = row.district
        cells[1].text = f"{row.double_exclusion_among_poor:.1%}"
        cells[2].text = f"{row.estimated_double_excluded_poor_households:,.0f}"
    format_table(table, widths=[2.2, 2.2, 2.35], font_size=9)

    add_para(doc,
        "Ngoma has the highest district rate (29.6%); Kamonyi has the largest estimated count (about 12,200). Nyanza is also high on both measures. Districts can rank differently by rate and by number of households, so both measures matter for planning.",
        size=9.4, after=1)

    doc.add_page_break()
    p = doc.add_paragraph(style="Title")
    set_paragraph(p, after=3)
    r = p.add_run("What decision makers can do")
    style_run(r, size=21, bold=True, color="000000")

    add_para(doc,
        "Use these district estimates to select places for service review, then verify the causes locally before changing eligibility or allocating individual benefits.",
        size=10.2, after=5)

    recommendations = [
        ("Review high-burden districts first.", "Start with Kamonyi, Nyagatare, Nyanza, Gisagara and Ngoma. Bring district social-protection teams, SACCOs, banks, mobile-money providers and civil-society groups together to check whether the mapped overlap reflects access barriers on the ground."),
        ("Offer an eligibility check and referral together.", "At local service points, let households ask for a VUP eligibility assessment and receive information about bank or mobile-money account options. The survey does not measure eligibility, so the assessment must precede any claim that a household was wrongly excluded."),
        ("Make account access practical and voluntary.", "Ask households what prevents account use, including fees, documentation, distance, phone access and confidence. Test accessible onboarding and support, with a clear choice of provider and payment method."),
        ("Track both reach and access.", "Monitor VUP participation, bank/mobile account access and the overlap among poor households. Pair administrative program data with periodic household evidence and publish district summaries with uncertainty intervals."),
    ]
    for title, body in recommendations:
        p = doc.add_paragraph(style="List Bullet")
        set_paragraph(p, after=5, line=1.04)
        r = p.add_run(title + " ")
        style_run(r, size=10, bold=True, color=NAVY)
        r = p.add_run(body)
        style_run(r, size=10)

    add_heading(doc, "Evidence and limits")
    add_para(doc,
        "The analysis uses the EICV7 cross-section (October 2023 to October 2024). District poverty rates use the population weight and reproduce the NISR Main Indicators Report's published district rates to 0.1 percentage point. Coverage and overlap estimates use household weights. Financial access means at least one reported bank or mobile-money account; tontine participation is shown separately. VUP coverage combines positive flags across Direct Support, Classic Public Works, Expanded Public Works, Nutrition-Sensitive Direct Support and the VUP financial-services module.",
        size=9.3, after=4)
    add_para(doc,
        "The estimates are descriptive, self-reported and cross-sectional. They cannot establish impact, program eligibility, barriers to access, or whether financial services are affordable or actively used. The 95% intervals use a stratified cluster linearization with a normal critical value; confirm the variance specification against NISR's survey-design documentation before publication. Nyarugenge and Kicukiro each have fewer than 50 sampled poor households and are flagged for caution.",
        size=9.3, after=4)

    add_heading(doc, "Sources")
    add_para(doc, "NISR, EICV7 Main Indicators Report 2023/24; NISR, EICV7 microdata catalog and questionnaire; authors' weighted calculations from the EICV7 cross-section. Fieldwork dates: October 2023-October 2024.", size=8.4, after=0, color=GRAY)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
